"""Normalize the World Bank Carbon Pricing Dashboard XLSX into wb_cpd.json.

License: CC-BY 4.0 (WB Data Catalog dataset 0042051). Attribution is emitted by the
export layer. Column contract pinned below per design/04 F2 - a drifted sheet fails
loudly here rather than mis-parsing.

Input:  corpus/directories/wb_cpd_<edition>.xlsx (fetched manually or by watcher alert)
Output: corpus/directories/wb_cpd.json
Usage:  py pipeline/harvest_wb.py corpus/directories/wb_cpd_2026-05.xlsx
"""
import datetime
import json
import re
import sys
import openpyxl
from common import ROOT

OUT = ROOT / "corpus" / "directories" / "wb_cpd.json"

SHEET = "Compliance_Gen Info"
HEADER_ROW = 5  # 1-based
EXPECTED = {0: "Unique ID", 1: "Instrument name", 2: "Type", 3: "Status",
            4: "Jurisdiction covered", 38: "Description", 39: "Recent developments",
            43: "Relation to other instruments"}

TYPE_MAP = {"Carbon tax": "carbon_tax", "ETS": "ets"}
STATUS_MAP = {"Implemented": "in_force", "Under development": "under_development",
              "Under consideration": "under_consideration", "Abolished": "abolished"}

# WB unique id -> our curated instrument id (aliases keep curated rows authoritative)
ALIASES = {"ETS_EU": "eu-ets", "ETS_TR": "tr-ets"}


def slug(wb_id, wb_type):
    """Tax_AL -> al-carbon-tax, ETS_US_WA -> us-wa-ets."""
    parts = wb_id.split("_")[1:]
    code = "-".join(p.lower() for p in parts)
    kind = "ets" if wb_type == "ETS" else "carbon-tax"
    return f"{code}-{kind}"


def main():
    xlsx = sys.argv[1] if len(sys.argv) > 1 else str(
        sorted((ROOT / "corpus" / "directories").glob("wb_cpd_*.xlsx"))[-1])
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    ws = wb[SHEET]
    rows = ws.iter_rows(min_row=HEADER_ROW, values_only=True)
    header = next(rows)
    for idx, want in EXPECTED.items():
        got = str(header[idx] or "").strip()
        if got != want:
            raise SystemExit(f"COLUMN CONTRACT BROKEN: col {idx} is {got!r}, expected {want!r}. "
                             "Inspect the sheet and re-pin (design/04 F2).")
    note_line = str(next(ws.iter_rows(min_row=1, max_row=1, values_only=True))[0] or "")
    instruments, skipped = [], []
    for r in rows:
        wb_id = r[0]
        if not wb_id:
            continue
        wb_type = str(r[2] or "").strip()
        if wb_type not in TYPE_MAP:
            skipped.append((wb_id, wb_type))
            continue
        wb_status = str(r[3] or "").strip()
        instruments.append({
            "wb_id": wb_id,
            "our_id": ALIASES.get(wb_id, slug(wb_id, wb_type)),
            "name": str(r[1] or "").strip(),
            "instrument_type": TYPE_MAP[wb_type],
            "wb_status": wb_status,
            "status": STATUS_MAP[wb_status],
            "jurisdiction_label": str(r[4] or "").strip(),
            "jurisdiction_code": "-".join(p.upper() for p in wb_id.split("_")[1:]),
            "description": (str(r[38]).strip() if r[38] else None),
            "relation_note": (str(r[43]).strip() if r[43] else None),
        })
    result = {"run": datetime.date.today().isoformat(),
              "edition": xlsx.split("wb_cpd_")[-1].replace(".xlsx", ""),
              "data_note": note_line,
              "source_url": "https://carbonpricingdashboard.worldbank.org/about-us#download-data",
              "license": "cc-by-4.0",
              "instruments": instruments,
              "skipped": skipped}
    OUT.write_text(json.dumps(result, indent=1, ensure_ascii=False), encoding="utf-8")
    print(f"{len(instruments)} instruments normalized ({len(skipped)} skipped: {skipped}) -> {OUT}")


if __name__ == "__main__":
    main()
